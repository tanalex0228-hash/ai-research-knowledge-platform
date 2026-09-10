"""Column-driven student account importer for Django Admin.

Unlike the canonical StudentRoster importer, this utility deliberately does
not attach meaning to worksheet names, classes, academic years, or student
numbers.  It finds every worksheet with the required account columns and
creates only rows explicitly identified as students.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Iterable
from uuid import uuid4

from django.conf import settings
from django.contrib.auth.hashers import make_password
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from .models import Role, User, UserProfile, UserRole
from .services import record_audit_event, validate_student_email_domain


REQUIRED_COLUMNS = frozenset({"姓名", "帳號", "身份", "預設密碼", "權限群組"})


@dataclass(frozen=True)
class StudentAccountRow:
    sheet_name: str
    row_number: int
    display_name: str
    email: str
    password: str
    username: str

    @property
    def locator(self) -> str:
        return f"{self.sheet_name}!{self.row_number}"


@dataclass(frozen=True)
class StudentAccountImportPlan:
    rows: tuple[StudentAccountRow, ...]
    source: str
    source_version: str


class StudentAccountWorkbookError(ValidationError):
    """Raised before a generic student account workbook writes data."""


def _cell_text(value) -> str:
    return str(value).strip() if value is not None else ""


def _workbook_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _format_errors(errors: Iterable[str]) -> str:
    return "\n".join(f"- {error}" for error in errors)


def build_student_account_import_plan(
    path: str | Path,
    *,
    source: str | None = None,
) -> StudentAccountImportPlan:
    """Validate every worksheet containing the required student columns.

    Sheets without the account columns are ignored.  This means the workbook
    may use any sheet names and may contain unrelated worksheets.  Rows whose
    ``身份`` is not ``student`` are ignored rather than converted into another
    account type.
    """

    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise StudentAccountWorkbookError(f"Workbook does not exist: {workbook_path}")

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    except Exception as exc:
        raise StudentAccountWorkbookError("Unable to read the .xlsx workbook.") from exc

    rows: list[StudentAccountRow] = []
    errors: list[str] = []
    matched_sheets = 0
    for worksheet in workbook.worksheets:
        header = [_cell_text(value) for value in next(worksheet.iter_rows(values_only=True), ())]
        columns = {name: index for index, name in enumerate(header) if name}
        if not REQUIRED_COLUMNS.issubset(columns):
            continue
        matched_sheets += 1

        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            value_for = lambda column: (
                values[columns[column]] if columns[column] < len(values) else None
            )
            identity = _cell_text(value_for("身份")).casefold()
            display_name = _cell_text(value_for("姓名"))
            email = User.objects.normalize_email(_cell_text(value_for("帳號")))
            password = _cell_text(value_for("預設密碼"))
            group = _cell_text(value_for("權限群組")).casefold()

            # Empty template rows often retain spreadsheet formulas in other
            # columns.  Only an explicit student identity makes a row importable.
            if not identity:
                continue
            if identity != "student":
                continue
            if group != "student":
                errors.append(f"{worksheet.title}!{row_number}: 權限群組 must be student")
                continue
            if not all((display_name, email, password)):
                errors.append(
                    f"{worksheet.title}!{row_number}: 姓名、帳號與預設密碼 are required"
                )
                continue
            try:
                validate_student_email_domain(email)
            except ValidationError:
                errors.append(
                    f"{worksheet.title}!{row_number}: 帳號 must use an allowed student email domain"
                )
                continue
            rows.append(
                StudentAccountRow(
                    sheet_name=worksheet.title,
                    row_number=row_number,
                    display_name=display_name,
                    email=email,
                    password=password,
                    username=display_name,
                )
            )

    if not matched_sheets:
        errors.append("No worksheet contains the required student account columns")

    for key_name, key in (("帳號", lambda row: row.email.casefold()),):
        grouped: dict[str, list[StudentAccountRow]] = {}
        for row in rows:
            grouped.setdefault(key(row), []).append(row)
        for duplicate_rows in grouped.values():
            if len(duplicate_rows) > 1:
                errors.append(
                    f"Duplicate {key_name}: {', '.join(row.locator for row in duplicate_rows)}"
                )

    names: dict[str, list[StudentAccountRow]] = {}
    for row in rows:
        names.setdefault(row.display_name.casefold(), []).append(row)
    rows = [
        replace(row, username=row.email)
        if len(names[row.display_name.casefold()]) > 1
        else row
        for row in rows
    ]

    if not rows:
        errors.append("No student rows found in worksheets with the required columns")
    if errors:
        raise StudentAccountWorkbookError(_format_errors(errors))

    return StudentAccountImportPlan(
        rows=tuple(rows),
        source=source or workbook_path.name,
        source_version=_workbook_sha256(workbook_path),
    )


def apply_student_account_import_plan(
    plan: StudentAccountImportPlan,
    *,
    actor: User | None = None,
    request_id: str = "",
) -> dict[str, int]:
    """Create student User/Profile/Role/Group records without roster records.

    Existing accounts are deliberately left intact, including their passwords.
    No class, academic-year, StudentRoster, Professor, or teacher record is
    created or changed by this operation.
    """

    with transaction.atomic():
        student_role = (
            Role.objects.select_for_update().filter(slug="student", is_active=True).first()
        )
        if student_role is None:
            raise StudentAccountWorkbookError("The active student role does not exist")
        student_group, _ = Group.objects.get_or_create(name="student")

        rows = list(plan.rows)
        existing_users = {
            user.email: user
            for user in User.objects.filter(email__in=[row.email for row in rows])
        }
        username_owners = {
            user.username.casefold(): user.email
            for user in User.objects.filter(username__in=[row.username for row in rows])
        }
        # A new student with the same displayed name as an existing account
        # remains importable: its internal username becomes its school email.
        rows = [
            replace(row, username=row.email)
            if (
                row.email not in existing_users
                and username_owners.get(row.username.casefold()) not in (None, row.email)
            )
            else row
            for row in rows
        ]
        username_owners = {
            user.username.casefold(): user.email
            for user in User.objects.filter(username__in=[row.username for row in rows])
        }
        errors = [
            f"{row.locator}: 使用者名稱 already belongs to a different account"
            for row in rows
            if row.email not in existing_users
            and username_owners.get(row.username.casefold()) not in (None, row.email)
        ]
        if errors:
            raise StudentAccountWorkbookError(_format_errors(errors))

        results = {"created_users": 0, "existing_users": 0, "created_profiles": 0}
        audit_request_id = request_id or f"student-user-import:{uuid4()}"
        for row in rows:
            user = existing_users.get(row.email)
            if user is not None:
                results["existing_users"] += 1
                continue

            user = User(username=row.username, email=row.email, is_active=True)
            user.password = make_password(row.password)
            user.save()
            UserRole.objects.create(user=user, role=student_role, assigned_by=actor)
            user.groups.add(student_group)
            UserProfile.objects.create(user=user, display_name=row.display_name)
            results["created_users"] += 1
            results["created_profiles"] += 1

        record_audit_event(
            event_type="accounts.student_user_import.completed",
            actor=actor,
            request_id=audit_request_id,
            metadata={
                "created_users": results["created_users"],
                "existing_users": results["existing_users"],
                "source": plan.source,
                "source_version": plan.source_version,
            },
        )
    return results
