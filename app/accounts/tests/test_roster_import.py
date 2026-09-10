from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path

from django.core.management import call_command
from django.core.management.base import CommandError
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook

from accounts.models import Role, StudentRoster, User, UserProfile, UserRole
from accounts.roster_import import (
    STUDENT_SHEET_NAMES,
    RosterWorkbookError,
    apply_import_plan,
    build_import_plan,
)
from accounts.student_user_import import (
    apply_student_account_import_plan,
    build_student_account_import_plan,
)
from professors.models import Professor


HEADERS = ("序號", "姓名", "帳號", "學號", "身份", "單位", "預設密碼", "權限群組")


class SystemRosterImportTests(TestCase):
    def _write_workbook(self, rows_by_sheet: dict[str, list[tuple]] | None = None) -> Path:
        workbook = Workbook()
        first = workbook.active
        first.title = "教師名單"
        first.append(HEADERS)
        for sheet_name in STUDENT_SHEET_NAMES:
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(HEADERS)
            for row in (rows_by_sheet or {}).get(sheet_name, []):
                sheet.append(row)
        temp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp.close()
        path = Path(temp.name)
        workbook.save(path)
        self.addCleanup(path.unlink, missing_ok=True)
        return path

    def _write_generic_user_workbook(self, sheets: dict[str, list[tuple]]) -> Path:
        workbook = Workbook()
        first = workbook.active
        first.title = next(iter(sheets))
        for index, (sheet_name, rows) in enumerate(sheets.items()):
            sheet = first if index == 0 else workbook.create_sheet(sheet_name)
            sheet.append(HEADERS)
            for row in rows:
                sheet.append(row)
        temp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        temp.close()
        path = Path(temp.name)
        workbook.save(path)
        self.addCleanup(path.unlink, missing_ok=True)
        return path

    @staticmethod
    def _student_row(*, name="王小明", student_id="414411001", password="temporary-password"):
        return (
            1,
            name,
            f"{student_id}@cloud.fju.edu.tw",
            student_id,
            "student",
            "金企二甲",
            password,
            "student",
        )

    def test_apply_creates_only_student_accounts_and_roster(self):
        path = self._write_workbook({"金企二甲": [self._student_row()]})

        plan = build_import_plan(path)
        result = apply_import_plan(plan, effective_start=date(2026, 9, 1))

        self.assertEqual(result["created_users"], 1)
        user = User.objects.get(email="414411001@cloud.fju.edu.tw")
        self.assertEqual(user.username, "王小明")
        self.assertTrue(user.check_password("temporary-password"))
        self.assertTrue(user.has_platform_role("student"))
        self.assertTrue(user.groups.filter(name="student").exists())
        self.assertFalse(user.is_staff)
        profile = UserProfile.objects.get(user=user)
        self.assertEqual(profile.display_name, "王小明")
        self.assertEqual(profile.roster_entry.academic_year, 2)
        self.assertFalse(profile.roster_entry.project_eligibility)
        self.assertEqual(Professor.objects.count(), 0)

    def test_preview_command_never_writes_database(self):
        path = self._write_workbook({"金企二甲": [self._student_row()]})

        call_command("import_system_roster", str(path))

        self.assertEqual(User.objects.count(), 0)
        self.assertEqual(StudentRoster.objects.count(), 0)

    def test_duplicate_names_use_email_as_internal_usernames(self):
        path = self._write_workbook(
            {
                "金企二甲": [self._student_row(name="同名學生", student_id="414411001")],
                "金企二乙": [self._student_row(name="同名學生", student_id="414412001")],
            }
        )

        plan = build_import_plan(path)
        usernames = {row.username for row in plan.rows}
        self.assertEqual(
            usernames,
            {"414411001@cloud.fju.edu.tw", "414412001@cloud.fju.edu.tw"},
        )

    def test_duplicate_student_id_rejects_entire_workbook(self):
        path = self._write_workbook(
            {
                "金企三乙": [self._student_row(name="甲同學", student_id="413412093")],
                "金企四乙": [self._student_row(name="乙同學", student_id="413412093")],
            }
        )

        with self.assertRaises(RosterWorkbookError):
            build_import_plan(path)
        self.assertEqual(User.objects.count(), 0)

    def test_existing_account_password_is_not_reset(self):
        student_role = Role.objects.get(slug="student")
        user = User.objects.create_user(
            username="王小明",
            email="414411001@cloud.fju.edu.tw",
            password="existing-password",
        )
        UserRole.objects.create(user=user, role=student_role)
        path = self._write_workbook({"金企二甲": [self._student_row(password="new-password")]})

        result = apply_import_plan(build_import_plan(path))

        user.refresh_from_db()
        self.assertEqual(result["existing_users"], 1)
        self.assertTrue(user.check_password("existing-password"))
        self.assertFalse(user.check_password("new-password"))

    def test_reapplying_the_same_plan_is_idempotent(self):
        path = self._write_workbook({"金企二甲": [self._student_row()]})
        plan = build_import_plan(path)

        apply_import_plan(plan)
        result = apply_import_plan(plan)

        self.assertEqual(result["created_users"], 0)
        self.assertEqual(result["existing_users"], 1)
        self.assertEqual(User.objects.count(), 1)
        self.assertEqual(StudentRoster.objects.count(), 1)

    def test_teacher_sheet_is_not_required_or_read(self):
        path = self._write_workbook({"金企二甲": [self._student_row()]})

        plan = build_import_plan(path)

        self.assertEqual(len(plan.rows), 1)

    def test_explicit_exclusion_resolves_duplicate_student_id(self):
        path = self._write_workbook(
            {
                "金企三乙": [self._student_row(name="羅銘鋒", student_id="412412093")],
                "金企四乙": [self._student_row(name="羅銘鋒", student_id="412412093")],
            }
        )

        plan = build_import_plan(path, excluded_locations=("金企三乙!2",))

        self.assertEqual(len(plan.rows), 1)
        self.assertEqual(plan.rows[0].sheet_name, "金企四乙")

    def test_command_apply_stops_for_invalid_workbook(self):
        path = self._write_workbook({"金企二甲": [(1, "同學", "bad@example.com", "414411001", "student", "金企二甲", "password", "student")]})

        with self.assertRaises(CommandError):
            call_command("import_system_roster", str(path), "--apply")
        self.assertEqual(User.objects.count(), 0)

    def test_platform_admin_can_import_students_from_user_admin_page(self):
        administrator = User.objects.create_superuser(
            username="roster-admin",
            email="roster-admin@example.test",
            password="admin-password",
        )
        path = self._write_generic_user_workbook(
            {"任意新增工作表": [self._student_row(password="new-student-password")]}
        )
        self.client.force_login(administrator)

        page = self.client.get("/admin/accounts/user/import-students/")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "批量匯入學生帳號")

        response = self.client.post(
            "/admin/accounts/user/import-students/",
            {
                "workbook": SimpleUploadedFile(
                    "系統名單.xlsx",
                    path.read_bytes(),
                    content_type=(
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ),
                ),
                "confirm": "on",
            },
        )

        self.assertRedirects(response, "/admin/accounts/user/")
        user = User.objects.get(email="414411001@cloud.fju.edu.tw")
        self.assertTrue(user.check_password("new-student-password"))
        self.assertTrue(user.has_platform_role("student"))
        self.assertTrue(user.groups.filter(name="student").exists())
        self.assertEqual(StudentRoster.objects.count(), 0)
        self.assertEqual(Professor.objects.count(), 0)

    def test_generic_user_import_scans_any_sheet_and_ignores_non_students(self):
        path = self._write_generic_user_workbook(
            {
                "九月新增": [self._student_row()],
                "其他人員": [
                    (
                        1,
                        "教師甲",
                        "teacher@example.edu.tw",
                        "T001",
                        "teacher",
                        "任意單位",
                        "not-used",
                        "teacher",
                    )
                ],
            }
        )

        plan = build_student_account_import_plan(path)
        result = apply_student_account_import_plan(plan)

        self.assertEqual(len(plan.rows), 1)
        self.assertEqual(plan.rows[0].sheet_name, "九月新增")
        self.assertEqual(result["created_users"], 1)
        self.assertEqual(StudentRoster.objects.count(), 0)
        self.assertFalse(User.objects.filter(email="teacher@example.edu.tw").exists())
