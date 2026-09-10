"""Strict importer for the department's fixed-format student roster workbook.

The importer deliberately knows nothing about professors.  A workbook is first
fully validated and only then may its student accounts and roster records be
created in one database transaction.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, replace
from datetime import date
from pathlib import Path
from typing import Iterable
from uuid import uuid4

from django.conf import settings
from django.contrib.auth.models import Group
from django.contrib.auth.hashers import make_password
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    EnrollmentStatus,
    Program,
    Role,
    StudentRoster,
    User,
    UserProfile,
    UserRole,
)
from .services import record_audit_event


STUDENT_SHEET_NAMES = (
    "金企二甲",
    "金企二乙",
    "金企三甲",
    "金企三乙",
    "金企四甲",
    "金企四乙",
)

REQUIRED_COLUMNS = frozenset({"姓名", "帳號", "學號", "身份", "預設密碼", "權限群組"})


@dataclass(frozen=True)
class RosterRow:
    sheet_name: str
    row_number: int
    display_name: str
    email: str
    student_id: str
    password: str
    academic_year: int
    username: str = ""

    @property
    def locator(self) -> str:
        return f"{self.sheet_name}!{self.row_number}"


@dataclass(frozen=True)
class ImportPlan:
    rows: tuple[RosterRow, ...]
    source: str
    source_version: str


class RosterWorkbookError(ValidationError):
    """Raised before any database write when the workbook is not trustworthy."""


def _cell_text(value) -> str:
    return str(value).strip() if value is not None else ""


def _student_id(value) -> str:
    # Excel may deserialize an all-numeric student number as an integer.
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _cell_text(value)


def _academic_year_for_sheet(sheet_name: str) -> int:
    academic_years = {
        "金企二甲": 2,
        "金企二乙": 2,
        "金企三甲": 3,
        "金企三乙": 3,
        "金企四甲": 4,
        "金企四乙": 4,
    }
    try:
        return academic_years[sheet_name]
    except KeyError as exc:
        raise RosterWorkbookError(f"Unsupported student sheet name: {sheet_name}") from exc


def workbook_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _format_errors(errors: Iterable[str]) -> str:
    return "\n".join(f"- {error}" for error in errors)


def build_import_plan(
    path: str | Path,
    *,
    source: str | None = None,
    excluded_locations: Iterable[str] = (),
) -> ImportPlan:
    """Read and validate the fixed workbook format without touching the database."""

    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise RosterWorkbookError(f"Workbook does not exist: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    missing_sheets = [name for name in STUDENT_SHEET_NAMES if name not in workbook.sheetnames]
    if missing_sheets:
        raise RosterWorkbookError(f"Missing required sheets: {', '.join(missing_sheets)}")

    excluded = {location.strip() for location in excluded_locations if location.strip()}
    rows: list[RosterRow] = []
    errors: list[str] = []
    for sheet_name in STUDENT_SHEET_NAMES:
        worksheet = workbook[sheet_name]
        header = [_cell_text(value) for value in next(worksheet.iter_rows(values_only=True), ())]
        columns = {name: index for index, name in enumerate(header) if name}
        missing_columns = sorted(REQUIRED_COLUMNS - columns.keys())
        if missing_columns:
            errors.append(f"{sheet_name}: missing columns {', '.join(missing_columns)}")
            continue

        for row_number, values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            locator = f"{sheet_name}!{row_number}"
            if locator in excluded:
                continue
            value_for = lambda column: values[columns[column]] if columns[column] < len(values) else None
            identity = _cell_text(value_for("身份")).casefold()
            display_name = _cell_text(value_for("姓名"))
            student_id = _student_id(value_for("學號"))
            email = User.objects.normalize_email(_cell_text(value_for("帳號")))
            password = _cell_text(value_for("預設密碼"))
            role = _cell_text(value_for("權限群組")).casefold()

            # The supplied workbook contains formatted but otherwise empty rows.
            # Excel templates prefill formula-like account/password cells even
            # when there is no person in the row.  A row without any identity
            # value is therefore not a candidate account and is ignored.
            if not any((identity, display_name, student_id)):
                continue
            if identity != "student":
                errors.append(f"{sheet_name}!{row_number}: 身份 must be student")
                continue
            if role != "student":
                errors.append(f"{sheet_name}!{row_number}: 權限群組 must be student")
                continue
            if not all((display_name, student_id, email, password)):
                errors.append(f"{sheet_name}!{row_number}: 姓名、帳號、學號與預設密碼 are required")
                continue
            expected_email = f"{student_id}@cloud.fju.edu.tw"
            if email != expected_email:
                errors.append(f"{sheet_name}!{row_number}: 帳號 must match 學號 and the FJU cloud domain")
                continue
            rows.append(
                RosterRow(
                    sheet_name=sheet_name,
                    row_number=row_number,
                    display_name=display_name,
                    email=email,
                    student_id=student_id,
                    password=password,
                    academic_year=_academic_year_for_sheet(sheet_name),
                    username=display_name,
                )
            )

    for key_name, key in (("帳號", lambda row: row.email.casefold()), ("學號", lambda row: row.student_id.casefold())):
        grouped: dict[str, list[RosterRow]] = {}
        for row in rows:
            grouped.setdefault(key(row), []).append(row)
        for duplicate_rows in grouped.values():
            if len(duplicate_rows) > 1:
                locations = ", ".join(row.locator for row in duplicate_rows)
                errors.append(f"Duplicate {key_name}: {locations}")

    # The platform keeps the supplied name as the public profile name.  Django
    # usernames must be unique, so only duplicate names use their already
    # verified FJU cloud email as the internal username.
    names: dict[str, list[RosterRow]] = {}
    for row in rows:
        names.setdefault(row.display_name.casefold(), []).append(row)
    rows = [
        replace(row, username=row.email) if len(names[row.display_name.casefold()]) > 1 else row
        for row in rows
    ]

    if not rows:
        errors.append("No student rows found in the required sheets")
    if errors:
        raise RosterWorkbookError(_format_errors(errors))

    return ImportPlan(
        rows=tuple(rows),
        source=source or workbook_path.name,
        source_version=workbook_sha256(workbook_path),
    )


def apply_import_plan(
    plan: ImportPlan,
    *,
    actor: User | None = None,
    effective_start: date | None = None,
) -> dict[str, int]:
    """Create only student identities, their roster records, and student roles.

    Existing accounts retain their password.  An account with a conflicting
    identity is rejected before any writes.  No professor model is queried or
    changed by this operation.
    """

    effective_start = effective_start or timezone.localdate()
    rows = plan.rows
    emails = [row.email for row in rows]
    usernames = [row.username for row in rows]
    student_ids = [row.student_id for row in rows]
    with transaction.atomic():
        # Lock the one system student role before reading identities.  This is a
        # stable mutex shared by every roster import, so two operators cannot
        # both pass preflight against the same stale account snapshot.
        student_role = (
            Role.objects.select_for_update().filter(slug="student", is_active=True).first()
        )
        if student_role is None:
            raise RosterWorkbookError("The active student role does not exist")
        # Django groups are used by the administration permission picker.  Keep
        # this in sync with the normalized platform role so a roster-created
        # student has both the public-site identity and the configured
        # administration permission group.
        student_group, _ = Group.objects.get_or_create(name="student")
        existing_users = {
            user.email: user
            for user in User.objects.filter(email__in=emails).prefetch_related("roles")
        }
        existing_usernames = {
            user.username.casefold(): user.email
            for user in User.objects.filter(username__in=usernames)
        }
        existing_rosters = {
            (roster.student_id, roster.fju_cloud_email): roster
            for roster in StudentRoster.objects.filter(
                student_id__in=student_ids,
                fju_cloud_email__in=emails,
                source_version=plan.source_version,
            )
        }

        errors: list[str] = []
        for row in rows:
            user_with_username = existing_usernames.get(row.username.casefold())
            if user_with_username and user_with_username != row.email:
                errors.append(f"{row.locator}: 使用者名稱 already belongs to a different account")
            existing_user = existing_users.get(row.email)
            if existing_user and existing_user.username != row.username:
                errors.append(f"{row.locator}: 帳號 already exists with a different 使用者名稱")
            if existing_user and not existing_user.has_platform_role("student"):
                errors.append(f"{row.locator}: 帳號 already exists without the student role")
        if errors:
            raise RosterWorkbookError(_format_errors(errors))

        results = {
            "created_users": 0,
            "existing_users": 0,
            "created_rosters": 0,
            "updated_profiles": 0,
        }
        request_id = f"roster-import:{uuid4()}"
        for row in rows:
            roster = existing_rosters.get((row.student_id, row.email))
            if roster is None:
                roster = StudentRoster.objects.create(
                    student_id=row.student_id,
                    fju_cloud_email=row.email,
                    department_code=settings.PRIMARY_DEPARTMENT_CODE,
                    program=Program.UNDERGRADUATE,
                    academic_year=row.academic_year,
                    enrollment_status=EnrollmentStatus.ACTIVE,
                    # Eligibility must be granted through the separate project
                    # policy; importing a class roster never grants upload rights.
                    project_eligibility=False,
                    effective_start=effective_start,
                    source=plan.source,
                    source_version=plan.source_version,
                    imported_by=actor,
                )
                results["created_rosters"] += 1

            user = existing_users.get(row.email)
            if user is None:
                user = User(
                    username=row.username,
                    email=row.email,
                    is_active=True,
                )
                user.password = make_password(row.password)
                user.save()
                UserRole.objects.create(user=user, role=student_role, assigned_by=actor)
                results["created_users"] += 1
                record_audit_event(
                    event_type="accounts.student_roster_import.created_user",
                    actor=actor,
                    target_type="accounts.User",
                    target_id=user.id,
                    request_id=request_id,
                    metadata={"student_id": row.student_id, "source_version": plan.source_version},
                )
            else:
                results["existing_users"] += 1
                UserRole.objects.get_or_create(
                    user=user,
                    role=student_role,
                    defaults={"assigned_by": actor},
                )
            user.groups.add(student_group)

            profile, created = UserProfile.objects.get_or_create(
                user=user,
                defaults={"roster_entry": roster, "display_name": row.display_name},
            )
            if not created and (
                profile.roster_entry_id != roster.id or profile.display_name != row.display_name
            ):
                profile.roster_entry = roster
                profile.display_name = row.display_name
                profile.save(update_fields=("roster_entry", "display_name", "updated_at"))
                results["updated_profiles"] += 1

        record_audit_event(
            event_type="accounts.student_roster_import.completed",
            actor=actor,
            request_id=request_id,
            metadata={
                "created_users": results["created_users"],
                "existing_users": results["existing_users"],
                "created_rosters": results["created_rosters"],
                "source": plan.source,
                "source_version": plan.source_version,
            },
        )
    return results
